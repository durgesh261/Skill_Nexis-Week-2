"""
Week 2: SQL for Data Analysis Execution & Report Generator
Executes SQL queries on SQLite database and exports results to Excel.
"""

import os
import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def run_sql_analysis():
    print("=" * 80)
    print("WEEK 2: SQL FOR DATA ANALYSIS — QUERY EXECUTION & EXCEL EXPORT")
    print("=" * 80)
    
    script_dir = os.path.dirname(os.path.abspath(__file__))
    db_path = os.path.join(script_dir, "..", "data", "awesome_chocolates.db")
    excel_output = os.path.join(script_dir, "Week2_SQL_Query_Results.xlsx")
    
    conn = sqlite3.connect(db_path)
    
    queries = {
        "Top_Salespersons": """
            SELECT 
                p.Salesperson,
                p.Team,
                p.Location,
                COUNT(*) AS Total_Orders,
                SUM(s.Boxes) AS Total_Boxes_Sold,
                ROUND(SUM(s.Amount), 2) AS Total_Revenue,
                ROUND(AVG(s.Amount), 2) AS Average_Order_Value,
                ROUND(SUM(s.Amount) / SUM(s.Boxes), 2) AS Revenue_Per_Box
            FROM sales s
            JOIN people p ON s.SPID = p.SPID
            GROUP BY p.Salesperson, p.Team, p.Location
            ORDER BY Total_Revenue DESC
            LIMIT 10;
        """,
        "AOV_By_Country": """
            SELECT 
                g.Geo AS Country,
                g.Region,
                COUNT(*) AS Total_Shipments,
                ROUND(SUM(s.Amount), 2) AS Total_Revenue,
                ROUND(AVG(s.Amount), 2) AS Average_Order_Value
            FROM sales s
            JOIN geo g ON s.GeoID = g.GeoID
            GROUP BY g.Geo, g.Region
            ORDER BY Total_Revenue DESC;
        """,
        "AOV_By_Category": """
            SELECT 
                pr.Category,
                COUNT(*) AS Total_Orders,
                ROUND(SUM(s.Amount), 2) AS Total_Revenue,
                ROUND(AVG(s.Amount), 2) AS Average_Order_Value,
                ROUND(SUM(s.Amount) * 100.0 / (SELECT SUM(Amount) FROM sales), 2) AS Revenue_Share_Pct
            FROM sales s
            JOIN products pr ON s.PID = pr.PID
            GROUP BY pr.Category
            ORDER BY Total_Revenue DESC;
        """,
        "Amount_Tiers_CASE": """
            SELECT 
                CASE 
                    WHEN Amount < 1000 THEN 'Under 1k'
                    WHEN Amount < 5000 THEN 'Under 5k'
                    WHEN Amount < 10000 THEN 'Under 10k'
                    ELSE '10k or more'
                END AS Amount_Category,
                COUNT(*) AS Total_Shipments,
                ROUND(SUM(Amount), 2) AS Total_Revenue,
                ROUND(AVG(Amount), 2) AS Average_Order_Value
            FROM sales
            GROUP BY Amount_Category
            ORDER BY Total_Revenue DESC;
        """,
        "P3_MilkBars_vs_Eclairs": """
            SELECT 
                pr.Product, 
                SUM(s.Boxes) AS Total_Boxes_Sold,
                ROUND(SUM(s.Amount), 2) AS Total_Revenue,
                COUNT(*) AS Total_Orders
            FROM sales s
            JOIN products pr ON s.PID = pr.PID
            WHERE pr.Product IN ('Milk Bars', 'Eclairs')
            GROUP BY pr.Product
            ORDER BY Total_Boxes_Sold DESC;
        """,
        "P8_Monthly_1k_Boxes": """
            SELECT 
                strftime('%Y-%m', SaleDate) AS Year_Month,
                COUNT(*) AS Shipments_Over_1000_Boxes,
                ROUND(SUM(Amount), 2) AS Total_Revenue
            FROM sales
            WHERE Boxes > 1000
            GROUP BY strftime('%Y-%m', SaleDate)
            ORDER BY Year_Month ASC;
        """,
        "P10_India_vs_Australia": """
            SELECT 
                strftime('%Y-%m', s.SaleDate) AS Year_Month,
                SUM(CASE WHEN g.Geo = 'India' THEN s.Boxes ELSE 0 END) AS India_Boxes,
                SUM(CASE WHEN g.Geo = 'Australia' THEN s.Boxes ELSE 0 END) AS Australia_Boxes,
                CASE 
                    WHEN SUM(CASE WHEN g.Geo = 'India' THEN s.Boxes ELSE 0 END) > SUM(CASE WHEN g.Geo = 'Australia' THEN s.Boxes ELSE 0 END) THEN 'India'
                    ELSE 'Australia'
                END AS Leading_Market
            FROM sales s
            JOIN geo g ON s.GeoID = g.GeoID
            GROUP BY strftime('%Y-%m', s.SaleDate)
            ORDER BY Year_Month ASC;
        """
    }
    
    results = {}
    for name, sql in queries.items():
        print(f"\n---> Executing Query: {name}")
        df_q = pd.read_sql_query(sql, conn)
        results[name] = df_q
        print(df_q.head(5).to_string(index=False))
        
    conn.close()
    
    print("\nWriting formatted results to Excel workbook...")
    with pd.ExcelWriter(excel_output, engine="openpyxl") as writer:
        for name, df_res in results.items():
            df_res.to_excel(writer, sheet_name=name, index=False)
            
    # Apply corporate formatting to Excel sheets
    wb = openpyxl.load_workbook(excel_output)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(left=Side(style="thin", color="D9D9D9"),
                         right=Side(style="thin", color="D9D9D9"),
                         top=Side(style="thin", color="D9D9D9"),
                         bottom=Side(style="thin", color="D9D9D9"))
                         
    for name in wb.sheetnames:
        ws = wb[name]
        ws.views.sheetView[0].showGridLines = True
        
        # Style headers
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Style data rows and auto-fit columns
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                cell.border = thin_border
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
            
    wb.save(excel_output)
    print(f"\nAll SQL Query results exported to: {excel_output}")
    print("=" * 80)
    print("SQL ANALYSIS EXECUTION COMPLETED SUCCESSFULLY!")
    print("=" * 80)

if __name__ == "__main__":
    run_sql_analysis()
